"""Multi-format export pipeline — Markdown, HTML, PDF, DOCX, JSON, BibTeX."""

import json
import logging
from abc import ABC, abstractmethod
from pathlib import Path

logger = logging.getLogger(__name__)


class BaseExporter(ABC):
    """Base class for report exporters."""

    @abstractmethod
    async def export(self, markdown: str, metadata: dict, output_path: Path) -> Path:
        """Export report to file. Returns output path."""
        pass


class MarkdownExporter(BaseExporter):
    """Passthrough markdown exporter."""

    async def export(self, markdown: str, metadata: dict, output_path: Path) -> Path:
        """Export report as markdown file."""
        logger.info("MarkdownExporter.export start: format=markdown output_path=%s chars=%d", output_path, len(markdown))
        md_path = output_path.with_suffix(".md")
        md_path.write_text(markdown, encoding="utf-8")
        logger.info("MarkdownExporter.export done: exported report to %s (%d bytes)", md_path, md_path.stat().st_size)
        return md_path


class HTMLExporter(BaseExporter):
    """Export to styled HTML."""

    async def export(self, markdown: str, metadata: dict, output_path: Path) -> Path:
        """Export report as styled HTML file."""
        logger.info("HTMLExporter.export start: format=html output_path=%s chars=%d", output_path, len(markdown))
        if not markdown.strip():
            logger.warning("HTMLExporter: markdown content is empty for %s", output_path)
        import markdown as md_lib

        html_content = md_lib.markdown(markdown, extensions=["tables", "fenced_code", "toc"])
        logger.debug("HTMLExporter: rendered html_content chars=%d", len(html_content))

        styled_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>{metadata.get('title', 'Research Report')}</title>
    <style>
        body {{ font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif;
               max-width: 800px; margin: 0 auto; padding: 2rem; line-height: 1.6; color: #333; }}
        h1 {{ color: #1a1a1a; border-bottom: 2px solid #e0e0e0; padding-bottom: 0.5rem; }}
        h2 {{ color: #2c3e50; margin-top: 2rem; }}
        blockquote {{ border-left: 4px solid #3498db; margin: 1rem 0; padding: 0.5rem 1rem; background: #f8f9fa; }}
        code {{ background: #f4f4f4; padding: 2px 6px; border-radius: 3px; }}
        table {{ border-collapse: collapse; width: 100%; margin: 1rem 0; }}
        th, td {{ border: 1px solid #ddd; padding: 8px; text-align: left; }}
        th {{ background: #f8f9fa; }}
    </style>
</head>
<body>
{html_content}
</body>
</html>"""

        html_path = output_path.with_suffix(".html")
        html_path.write_text(styled_html, encoding="utf-8")
        logger.info("HTMLExporter.export done: exported report to %s (%d bytes)", html_path, html_path.stat().st_size)
        return html_path


class PDFExporter(BaseExporter):
    """Export to PDF via weasyprint."""

    async def export(self, markdown: str, metadata: dict, output_path: Path) -> Path:
        """Export report as PDF via weasyprint."""
        logger.info("PDFExporter.export start: format=pdf output_path=%s chars=%d", output_path, len(markdown))
        html_exporter = HTMLExporter()
        html_path = await html_exporter.export(markdown, metadata, output_path.with_suffix(".tmp"))

        try:
            from weasyprint import HTML

            pdf_path = output_path.with_suffix(".pdf")
            logger.debug("PDFExporter: rendering PDF from intermediate html %s", html_path)
            HTML(filename=str(html_path)).write_pdf(str(pdf_path))
            logger.info(
                "PDFExporter.export done: exported report to %s (%d bytes)", pdf_path, pdf_path.stat().st_size
            )
            return pdf_path
        except Exception:
            logger.exception("PDFExporter.export failed for %s", output_path)
            raise
        finally:
            logger.debug("PDFExporter: cleaning up intermediate html %s", html_path)
            html_path.unlink(missing_ok=True)


class DOCXExporter(BaseExporter):
    """Export to DOCX via python-docx."""

    async def export(self, markdown: str, metadata: dict, output_path: Path) -> Path:
        """Export report as DOCX via python-docx."""
        logger.info("DOCXExporter.export start: format=docx output_path=%s chars=%d", output_path, len(markdown))
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()
        title = doc.add_heading(metadata.get("title", "Research Report"), 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        lines = markdown.split("\n")
        for line in lines:
            if line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.startswith("- "):
                doc.add_paragraph(line[2:], style="List Bullet")
            elif line.strip():
                doc.add_paragraph(line)

        logger.debug("DOCXExporter: processed %d markdown line(s)", len(lines))
        docx_path = output_path.with_suffix(".docx")
        doc.save(str(docx_path))
        logger.info("DOCXExporter.export done: exported report to %s (%d bytes)", docx_path, docx_path.stat().st_size)
        return docx_path


class JSONExporter(BaseExporter):
    """Export structured JSON with all metadata."""

    async def export(self, markdown: str, metadata: dict, output_path: Path) -> Path:
        """Export report as structured JSON file."""
        logger.info("JSONExporter.export start: format=json output_path=%s chars=%d", output_path, len(markdown))
        structured = {
            "report": {
                "markdown": markdown,
                "title": metadata.get("title", ""),
                "generated_at": metadata.get("generated_at", ""),
                "profile": metadata.get("profile", ""),
                "word_count": len(markdown.split()),
            },
            "metadata": metadata,
        }

        json_path = output_path.with_suffix(".json")
        json_path.write_text(json.dumps(structured, indent=2), encoding="utf-8")
        logger.info("JSONExporter.export done: exported report to %s (%d bytes)", json_path, json_path.stat().st_size)
        return json_path


class BibTeXExporter(BaseExporter):
    """Export bibliography as BibTeX file."""

    async def export(self, markdown: str, metadata: dict, output_path: Path) -> Path:
        """Export bibliography as BibTeX file."""
        logger.info("BibTeXExporter.export start: format=bibtex output_path=%s", output_path)
        from open_deep_research.citation import CitationFormatter

        sources = metadata.get("sources", [])
        if not sources:
            logger.warning("BibTeXExporter: no sources in metadata for %s", output_path)
        formatter = CitationFormatter("vanilla")
        bibtex_content = formatter.export_bibtex(sources)
        logger.debug("BibTeXExporter: formatted %d source(s) into bibtex", len(sources))

        bib_path = output_path.with_suffix(".bib")
        bib_path.write_text(bibtex_content, encoding="utf-8")
        logger.info("BibTeXExporter.export done: exported bibliography to %s (%d bytes)", bib_path, bib_path.stat().st_size)
        return bib_path


EXPORTER_MAP = {
    "markdown": MarkdownExporter,
    "html": HTMLExporter,
    "pdf": PDFExporter,
    "docx": DOCXExporter,
    "json": JSONExporter,
    "bibtex": BibTeXExporter,
}


def get_exporter(format_name: str) -> BaseExporter:
    """Get exporter instance by format name."""
    cls = EXPORTER_MAP.get(format_name)
    if cls is None:
        logger.warning("get_exporter: unsupported export format requested: %s", format_name)
        raise ValueError(f"Unknown export format: {format_name}")
    logger.debug("get_exporter: resolved format %s -> %s", format_name, cls.__name__)
    return cls()
